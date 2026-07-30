#!/usr/bin/env python3
"""Export docs/fm-future-storyboard.xlsx (and the Rev 1 version, for the redline)
to JSON for build-deck.js. Run from the repo root."""
import json, subprocess, sys, io
import openpyxl

REV1 = '6bb3491'  # last commit before the Jul 2026 feedback pass


def dump(ws):
    rows, t = [], 0
    tc = lambda s: f'{s // 60}:{s % 60:02d}'
    for r in range(2, 22):
        g = lambda c: str(ws.cell(r, c).value or '')
        d = int(ws.cell(r, 7).value)
        rows.append({'shot': int(ws.cell(r, 1).value), 'element': g(2), 'script': g(3),
                     'onscreen': g(4), 'visual': g(5), 'narrator': g(6), 'sec': d,
                     'vo': g(8), 'inTc': tc(t), 'outTc': tc(t + d)})
        t += d
    return rows, t, tc(t)


cur = openpyxl.load_workbook('docs/fm-future-storyboard.xlsx')['Sheet1']
rows, total, totalTc = dump(cur)
json.dump({'rows': rows, 'total': total, 'totalTc': totalTc}, open('sb.json', 'w'), indent=1)

blob = subprocess.run(['git', 'show', f'{REV1}:docs/fm-future-storyboard.xlsx'],
                      capture_output=True, check=True).stdout
old = openpyxl.load_workbook(io.BytesIO(blob))['Sheet1']
oldrows, _, _ = dump(old)
json.dump({str(r['shot']): {'script': r['script'], 'onscreen': r['onscreen'], 'sec': r['sec']}
           for r in oldrows}, open('sb-old.json', 'w'), indent=1)

print(f'sb.json + sb-old.json written. {len(rows)} shots, total {totalTc}.', file=sys.stderr)
